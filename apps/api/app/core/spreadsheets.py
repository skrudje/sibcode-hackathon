"""One bounded XLSX reader for staff uploads and the brand portal."""
import io
from contextlib import contextmanager
from pathlib import Path
from zipfile import BadZipFile, ZIP_DEFLATED, ZIP_STORED, ZipFile

from defusedxml.ElementTree import iterparse
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook

from .errors import AppError

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000
MAX_COLUMNS = 200
MAX_XML_BYTES = 16 * 1024 * 1024
MAX_UNPACKED_BYTES = 32 * 1024 * 1024


@contextmanager
def spreadsheet(value):
    content = value if isinstance(value, bytes) else Path(value).read_bytes()
    if len(content) > MAX_BYTES:
        raise AppError('file_too_large', 'Максимальный размер файла — 10 MiB', 413)
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if (len(entries) > 128 or len({entry.filename for entry in entries}) != len(entries)
                    or sum(entry.file_size for entry in entries) > MAX_UNPACKED_BYTES
                    or any(entry.file_size > MAX_XML_BYTES or entry.flag_bits & 1
                           or entry.compress_type not in {ZIP_STORED, ZIP_DEFLATED} for entry in entries)):
                raise AppError('file_too_large', 'Книга превышает безопасные ограничения XLSX', 413)
            for entry in entries:
                if entry.filename.endswith(('.xml', '.rels')):
                    with archive.open(entry) as stream:
                        for _, element in iterparse(stream, events=('end',), forbid_dtd=True,
                                                    forbid_entities=True, forbid_external=True):
                            element.clear()
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            sheet = book.active
            if not sheet or (sheet.max_column or 0) > MAX_COLUMNS:
                raise AppError('invalid_file_schema', 'В таблице допускается до 200 колонок')
            if (sheet.max_row or 0) > MAX_ROWS + 1:
                raise AppError('too_many_rows', 'В одном файле допускается до 10 000 строк данных')
            yield book
        finally:
            book.close()
    except AppError:
        raise
    except DefusedXmlException as exc:
        raise AppError('unsafe_spreadsheet', 'Книга содержит запрещённые XML-объявления') from exc
    except Exception as exc:
        raise AppError('invalid_file_schema', 'Не удалось прочитать книгу XLSX') from exc
