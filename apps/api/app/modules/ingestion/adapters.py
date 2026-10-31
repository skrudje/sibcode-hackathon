import csv
import json
from collections import Counter
from pathlib import Path

import polars as pl
from openpyxl import load_workbook

from app.core.errors import AppError

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000
REQUIRED_COLUMNS = {
    'contacts': ('phone', 'email', 'time'),
    'crm_csv': ('customer_id', 'event_id', 'event_type', 'event_time'),
    'pos_xlsx': ('id_клиента', 'id_чека', 'тип_события', 'дата_операции'),
}


def validate_headers(kind, headers):
    columns = [str(value).strip() if value is not None else '' for value in headers]
    missing = [name for name in REQUIRED_COLUMNS[kind] if name not in columns]
    repeated = [name for name, count in Counter(columns).items() if name and count > 1]
    empty = [str(index) for index, name in enumerate(columns, 1) if not name]
    problems, details = [], []
    if missing:
        problems.append('Отсутствуют обязательные колонки: ' + ', '.join(missing))
        details.extend({'field': name, 'message': 'Обязательная колонка отсутствует'} for name in missing)
    if repeated:
        # Header names can be user supplied; keep the job failure message bounded.
        names = ', '.join(name[:100] for name in repeated[:20])
        if len(repeated) > 20:
            names += ', …'
        problems.append('Повторяются колонки: ' + names)
        details.extend({'field': name[:100], 'message': 'Заголовок повторяется'} for name in repeated[:20])
    if empty:
        problems.append('Пустые заголовки в колонках: ' + ', '.join(empty[:20]))
    if problems:
        if kind == 'crm_csv' and len(columns) == 1 and ';' in columns[0]:
            problems.append('Используйте запятую как разделитель CSV')
        raise AppError('invalid_file_schema', '. '.join(problems) + '.', details=details)
    return columns


def read_rows(kind, path):
    try:
        if kind == 'contacts' and Path(path).suffix == '.json':
            value = json.loads(Path(path).read_text())
            rows = value['records']
            if not isinstance(rows, list):
                raise ValueError()
        elif kind == 'rest_api':
            value = json.loads(Path(path).read_text())
            rows = value['events']
            if not isinstance(rows, list):
                raise ValueError()
        elif kind == 'crm_csv' or (kind == 'contacts' and Path(path).suffix == '.csv'):
            # Inspect the original header before Polars renames duplicate names.
            with Path(path).open(encoding='utf-8-sig', newline='') as file:
                headers = next(csv.reader(file, strict=True), None)
            if headers is None:
                raise AppError('empty_import', 'CSV пустой: добавьте заголовки и строки событий')
            columns = validate_headers(kind, headers)
            frame = pl.read_csv(path, infer_schema=False, new_columns=columns)
            rows = frame.to_dicts()
        else:
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                iterator = book.active.iter_rows(values_only=True)
                headers = next(iterator, None)
                if headers is None:
                    raise AppError('empty_import', 'Активный лист XLSX пустой: добавьте заголовки и строки событий')
                headers = validate_headers(kind, headers)
                rows = []
                for row in iterator:
                    if len(rows) >= MAX_ROWS:
                        raise AppError('too_many_rows', 'В одном задании допускается до 10 000 строк')
                    rows.append(dict(zip(headers, row)))
            finally:
                book.close()
        if not rows and not (kind == 'contacts' and Path(path).suffix == '.json'):
            raise AppError('empty_import', 'Источник не содержит записей')
        if len(rows) > MAX_ROWS:
            raise AppError('too_many_rows', 'В одном задании допускается до 10 000 строк')
        return rows
    except AppError:
        raise
    except Exception as exc:
        if kind == 'crm_csv':
            raise AppError('invalid_file_schema', 'Не удалось прочитать CSV. Используйте UTF-8, разделитель-запятую и одинаковое число полей в строках.') from exc
        if kind == 'pos_xlsx':
            raise AppError('invalid_file_schema', 'Не удалось прочитать XLSX. Проверьте файл и таблицу с заголовками на активном листе.') from exc
        raise AppError('invalid_file_schema', 'Не удалось прочитать структуру источника. Проверьте шаблон.') from exc


def normalize(kind, row):
    if not isinstance(row, dict):
        raise AppError('missing_required_field', 'Запись должна быть объектом')
    if kind == 'contacts':
        from app.modules.integrations.service import normalized_event
        return normalized_event(row)
    if kind == 'rest_api':
        customer = row.get('customer') or {}
        if not isinstance(customer, dict):
            raise AppError('missing_required_field', 'Нет идентификатора клиента')
        return dict(external_customer_id=customer.get('id'), global_customer_key=customer.get('global_key'), display_name=customer.get('name'), email=customer.get('email'), phone=customer.get('phone'), push_token=customer.get('push_token'), external_event_id=row.get('id'), event_type=row.get('kind'), occurred_at=row.get('timestamp'), amount=row.get('total'), currency=row.get('currency'))
    if kind == 'crm_csv':
        keys = ['customer_id', 'global_customer_key', 'full_name', 'email', 'phone', 'event_id', 'event_type', 'event_time', 'amount', 'currency']
    else:
        keys = ['id_клиента', 'единый_ключ', 'имя', 'email', 'телефон', 'id_чека', 'тип_события', 'дата_операции', 'сумма', 'валюта']
    target = ['external_customer_id', 'global_customer_key', 'display_name', 'email', 'phone', 'external_event_id', 'event_type', 'occurred_at', 'amount', 'currency']
    return {to: (str(row.get(key)).strip() or None) if row.get(key) is not None else None for key, to in zip(keys, target)}
