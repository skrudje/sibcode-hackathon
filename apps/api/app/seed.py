import csv
import json
from datetime import datetime, timedelta
from uuid import UUID

from openpyxl import load_workbook
from pwdlib import PasswordHash
from sqlalchemy import select

from app.core.config import ROOT, Clock, settings
from app.core.db import SessionLocal, lock
from app.models import Campaign, ContactChannelState, Mall, RecommendationRule, Segment, Source, Tenant, User
from app.modules.radar.demo import generate_files as generate_radar_files


def generate_files(reference):
    template = ROOT / 'demo/templates'
    manifest = json.loads((template / 'demo_scenario.json').read_text())
    old = datetime.fromisoformat(manifest['reference_time'].replace('Z', '+00:00'))
    delta = reference - old
    out = settings.path('demo_dir')
    out.mkdir(parents=True, exist_ok=True)

    def shift(value):
        if not value:
            return value
        return (datetime.fromisoformat(value.replace('Z', '+00:00')) + delta).isoformat().replace('+00:00', 'Z')

    data = json.loads((template / 'api_A.json').read_text())
    for event in data['events']:
        event['timestamp'] = shift(event['timestamp'])
    (out / 'api_A.json').write_text(json.dumps(data, ensure_ascii=False))
    for name in ['crm_C_initial.csv', 'crm_C_after.csv', 'crm_C_conflict.csv']:
        with (template / name).open(encoding='utf-8-sig', newline='') as f:
            rows = list(csv.DictReader(f))
        for row in rows:
            row['event_time'] = shift(row['event_time'])
        with (out / name).open('w', encoding='utf-8-sig', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
    book = load_workbook(template / 'pos_B.xlsx')
    for row in book.active.iter_rows(min_row=2):
        row[7].value = shift(row[7].value)
    book.save(out / 'pos_B.xlsx')
    generate_radar_files(reference)
    return manifest


def seed():
    template = json.loads((ROOT / 'demo/templates/demo_scenario.json').read_text())
    ids = {k: UUID(v) for k, v in template['ids'].items()}
    with SessionLocal.begin() as db:
        lock(db, 'demo-seed')
        mall = db.get(Mall, ids['mall'])
        if not mall:
            mall = Mall(id=ids['mall'], name='Торговый центр «Орбита»', demo_reference_time=Clock.now())
            db.add(mall)
            db.flush()
        reference = mall.demo_reference_time
        manifest = generate_files(reference)
        if not db.scalar(select(User).where(User.email == settings.demo_email.lower())):
            db.add(User(mall_id=mall.id, email=settings.demo_email.lower(), password_hash=PasswordHash.recommended().hash(settings.demo_password)))
        for source_info, category in zip(manifest['sources'], ['Досуг', 'Еда и напитки', 'Книги']):
            tid, sid = UUID(source_info['tenant_id']), UUID(source_info['id'])
            if not db.get(Tenant, tid):
                db.add(Tenant(id=tid, mall_id=mall.id, code=source_info['tenant_code'], name=source_info['tenant_name'], category=category))
                db.flush()
            if not db.get(Source, sid):
                db.add(Source(id=sid, mall_id=mall.id, tenant_id=tid, name={'A':'API кинотеатра', 'B':'Касса кафе', 'C':'CRM книжного'}[source_info['tenant_code']], kind=source_info['kind'], covered_event_types=source_info['covered_event_types'], coverage_from=reference - timedelta(days=120), coverage_to=reference, config={'url': settings.mock_source_url} if source_info['kind'] == 'rest_api' else {}))
        if not db.get(Segment, ids['segment']):
            db.add(Segment(id=ids['segment'], mall_id=mall.id, name=manifest['segment']['name'], definition=manifest['segment']['definition']))
            db.flush()
        if not db.get(RecommendationRule, ids['recommendation_rule']):
            rule = {k: v for k, v in manifest['recommendation_rule'].items() if k not in ['valid_from_offset_days', 'valid_until_offset_days']}
            for key in ['id', 'segment_id', 'destination_tenant_id']:
                rule[key] = UUID(rule[key])
            db.add(RecommendationRule(mall_id=mall.id, **rule, valid_from=reference - timedelta(days=1), valid_until=reference + timedelta(days=30)))
            db.flush()
        if not db.get(Campaign, ids['campaign']):
            info = dict(manifest['campaign'])
            for key in ['id', 'segment_id', 'recommendation_rule_id']:
                info[key] = UUID(info[key])
            db.add(Campaign(mall_id=mall.id, **info))
        key = (mall.id, 'email', 'g008@example.invalid')
        if not db.get(ContactChannelState, key):
            db.add(ContactChannelState(mall_id=mall.id, channel='email', contact_normalized=key[2], last_exported_at=reference - timedelta(days=1)))
    settings.path('upload_dir').mkdir(parents=True, exist_ok=True)
    print('Demo catalog and files are ready. Customer events are loaded through imports.')


if __name__ == '__main__':
    seed()
